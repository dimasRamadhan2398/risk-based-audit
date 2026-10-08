import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

fpath = r'docs\ISO_27001\03_Checklist_Kesiapan_Bukti_93_Kontrol_Annex_A.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=False)
print('Sheet names:', wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n=======================================================")
    print(f"Sheet: {sname} (max_row={ws.max_row}, max_column={ws.max_column})")
    print(f"=======================================================")
    for r in range(1, min(ws.max_row + 1, 20)):
        row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 15))]
        if any(v is not None for v in row_vals):
            non_empty = [f"C{c}:{str(v)[:40]}" for c, v in enumerate(row_vals, 1) if v is not None]
            print(f"Row {r:2}: {non_empty}")
