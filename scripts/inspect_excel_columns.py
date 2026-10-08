import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

fpath = r'docs\ISO_27001\03_Checklist_Kesiapan_Bukti_93_Kontrol_Annex_A.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=False)
ws = wb['Checklist 93 Kontrol']

print("Columns in Row 6:")
for col_idx in range(1, ws.max_column + 1):
    val = ws.cell(6, col_idx).value
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    print(f"Col {col_idx:2} ({col_letter}): {val}")

print("\nFormulas in row 7:")
for col_idx in range(1, ws.max_column + 1):
    val = ws.cell(7, col_idx).value
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    if str(val).startswith('='):
        print(f"Col {col_letter}: {val}")

print("\nSample Row 7 data:")
for col_idx in range(1, ws.max_column + 1):
    val = ws.cell(7, col_idx).value
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    header = ws.cell(6, col_idx).value
    print(f"{col_letter} ({header}): {repr(val)}")
