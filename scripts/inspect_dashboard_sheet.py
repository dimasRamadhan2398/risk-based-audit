import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

fpath = r'docs\ISO_27001\03_Checklist_Kesiapan_Bukti_93_Kontrol_Annex_A.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=False)

ws_d = wb['Dashboard']
print("Dashboard Sheet:")
for r in range(1, 35):
    vals = [ws_d.cell(r, c).value for c in range(1, ws_d.max_column + 1)]
    if any(v is not None for v in vals):
        row_str = " | ".join([f"C{c}: {str(v)}" for c, v in enumerate(vals, 1) if v is not None])
        print(f"Row {r:2}: {row_str}")

ws_r = wb['Review Kesiapan']
print("\nReview Kesiapan Sheet (Header & Row 7):")
for r in [6, 7, 8]:
    vals = [ws_r.cell(r, c).value for c in range(1, ws_r.max_column + 1)]
    row_str = " | ".join([f"C{c}: {str(v)}" for c, v in enumerate(vals, 1) if v is not None])
    print(f"Row {r:2}: {row_str}")
