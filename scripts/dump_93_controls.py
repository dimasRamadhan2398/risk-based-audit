import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

fpath = r'docs\ISO_27001\03_Checklist_Kesiapan_Bukti_93_Kontrol_Annex_A.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=False)
ws = wb['Checklist 93 Kontrol']

with open('scratch_93_controls.txt', 'w', encoding='utf-8') as out:
    for r in range(7, 100):
        code = ws.cell(r, 2).value # Col B
        theme = ws.cell(r, 3).value # Col C
        title = ws.cell(r, 4).value # Col D
        focus = ws.cell(r, 5).value # Col E
        example = ws.cell(r, 6).value # Col F
        out.write(f"Row {r:2} | {code:6} | {theme:14} | {title} | {focus}\n")

print("Dumped 93 controls to scratch_93_controls.txt successfully!")
