import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

fpath = r'docs\ISO_27001\03_Checklist_Kesiapan_Bukti_93_Kontrol_Annex_A.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=False)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"Sheet: {sname} (max_row={ws.max_row}, max_column={ws.max_column})")

ws_main = wb['Checklist 93 Kontrol']
print("\n--- Header rows of 'Checklist 93 Kontrol' ---")
for r in range(1, 10):
    vals = [ws_main.cell(r, c).value for c in range(1, ws_main.max_column + 1)]
    if any(v is not None for v in vals):
        print(f"Row {r:2}: {[str(v)[:30] if v is not None else '' for v in vals[:26]]}")
