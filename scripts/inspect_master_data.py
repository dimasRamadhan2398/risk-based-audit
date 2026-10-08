import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

fpath = r'docs\ISO_27001\03_Checklist_Kesiapan_Bukti_93_Kontrol_Annex_A.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=False)
ws_m = wb['Master Data']

print("Master Data sheet content:")
for r in range(1, ws_m.max_row + 1):
    vals = [ws_m.cell(r, c).value for c in range(1, ws_m.max_column + 1)]
    if any(v is not None for v in vals):
        row_str = " | ".join([f"C{c}: {str(v)}" for c, v in enumerate(vals, 1) if v is not None])
        print(f"Row {r:2}: {row_str}")
